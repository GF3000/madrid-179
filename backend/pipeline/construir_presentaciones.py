"""Construye las presentaciones 1 (Descripción y problema) y 3 (Interfaz) sobre la plantilla de Slidesgo.
Reutiliza diapositivas de la plantilla (portada, índice, diagramas de ciclo, cierre) cambiando solo sus textos, y añade
diapositivas propias sobre su layout TITLE_ONLY con la paleta y tipografías del tema.
Salidas: docs/presentaciones/1_Descripcion_y_problema.pptx y docs/presentaciones/3_Interfaz.pptx
Después: powershell -File infra/scripts/pptx_a_pdf.ps1 <pptx> <pdf>   (requiere Poppins y Didact Gothic instaladas)
"""
import copy

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PLANTILLA = "docs/presentaciones/_plantilla/plantilla.pptx"
NAVY, MID, PERI, PALE = RGBColor(0x37, 0x47, 0x68), RGBColor(0x43, 0x56, 0x7F), RGBColor(0x9B, 0xAC, 0xCA), RGBColor(0xC2, 0xCD, 0xE5)
TINT, INK, GREY, WHITE = RGBColor(0xEE, 0xF1, 0xF8), RGBColor(0x1B, 0x1B, 0x1B), RGBColor(0x55, 0x5B, 0x68), RGBColor(0xFF, 0xFF, 0xFF)
HEAD, BODY = "Poppins", "Didact Gothic"
AUTORES = "Gregorio García Velasco (UC3M) · Guillermo Franco Gimeno (UPM)"


# ----------------------------------------------------------------------------------------------- utilidades
def por_id(slide, sid):
    def buscar(shapes):
        for s in shapes:
            if s.shape_id == sid:
                return s
            if s.shape_type == 6:
                r = buscar(s.shapes)
                if r is not None:
                    return r
    s = buscar(slide.shapes)
    assert s is not None, f"forma {sid} no encontrada"
    return s


def lineas(shape, textos, size=None, bold=None, color=None):
    """Sustituye el texto conservando el formato del primer párrafo y la primera ejecución de la plantilla."""
    tx = shape.text_frame._txBody
    ps = tx.findall(qn("a:p"))
    ppr = ps[0].find(qn("a:pPr"))
    r0 = ps[0].find(qn("a:r"))
    rpr = r0.find(qn("a:rPr")) if r0 is not None else None
    for p in ps:
        tx.remove(p)
    for t in textos:
        p = etree.SubElement(tx, qn("a:p"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        r = etree.SubElement(p, qn("a:r"))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        etree.SubElement(r, qn("a:t")).text = t
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
            if color is not None:
                r.font.color.rgb = color


def borrar(shape):
    shape._element.getparent().remove(shape._element)


def titulo(slide, principal, resto=""):
    t = slide.shapes.title
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = principal + (" " if resto else "")
    if resto:
        r2 = p.add_run()
        r2.text = resto
        r2.font.bold = False


def texto(slide, x, y, w, h, parrafos, size=12, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, espacio=4):
    """parrafos: lista de str o de listas de (texto, negrita[, color])."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, par in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(espacio)
        trozos = [(par, bold)] if isinstance(par, str) else par
        for tr in trozos:
            r = p.add_run()
            r.text = tr[0]
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = tr[1]
            r.font.color.rgb = tr[2] if len(tr) > 2 else color
    return tb


def caja(slide, x, y, w, h, fill=TINT, line=None, forma=MSO_SHAPE.ROUNDED_RECTANGLE, radio=0.08, dash=False):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radio
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
        if dash:
            s.line.dash_style = 4  # guiones
    s.shadow.inherit = False
    return s


def circulo(slide, cx, cy, d, txt, fill=NAVY, color=WHITE, size=14, font=HEAD):
    s = caja(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, forma=MSO_SHAPE.OVAL)
    tf = s.text_frame
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = txt
    r.font.size, r.font.name, r.font.bold, r.font.color.rgb = Pt(size), font, True, color
    return s


def tabla(slide, x, y, w, filas, anchos, size=11.5, alto_fila=0.36):
    n, m = len(filas), len(filas[0])
    t = slide.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(alto_fila * n)).table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(w * a)
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (TINT if i % 2 else WHITE)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.size, r.font.name = Pt(size), (HEAD if i == 0 else BODY)
            r.font.bold = i == 0 or j == 0
            r.font.color.rgb = WHITE if i == 0 else INK
            if j > 0 and isinstance(val, str) and val[:1] in "0123456789−+":
                p.alignment = PP_ALIGN.RIGHT
    return t


def flecha(slide, x1, y1, x2, y2, color=PERI):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def red_bayesiana_nativa(slide):
    """Los 8 nodos y 10 aristas de la red (red_bayesiana_viabilidad.py), dibujados con formas de la plantilla."""
    W_, H_ = 1.2, 0.48
    nodos = {  # centro (x, y); colocados para que ninguna arista atraviese otra caja
        "Talento": (5.0, 1.6), "Distancia Madrid": (7.45, 1.6), "Acceso ferroviario": (8.95, 1.6),
        "Renta": (4.25, 2.62), "Especialización": (5.6, 2.62), "Coste inmobiliario": (6.95, 2.62),
        "Dinamismo demográfico": (8.1, 3.64), "Viabilidad empresarial": (6.3, 4.66)}
    aristas = [("Talento", "Renta"), ("Talento", "Especialización"), ("Talento", "Coste inmobiliario"),
               ("Distancia Madrid", "Especialización"), ("Distancia Madrid", "Coste inmobiliario"),
               ("Distancia Madrid", "Dinamismo demográfico"), ("Distancia Madrid", "Acceso ferroviario"),
               ("Coste inmobiliario", "Dinamismo demográfico"), ("Especialización", "Viabilidad empresarial"),
               ("Dinamismo demográfico", "Viabilidad empresarial")]
    for a, b in aristas:  # primero las aristas, para que los nodos queden encima
        (xa, ya), (xb, yb) = nodos[a], nodos[b]
        if abs(ya - yb) < 0.01:  # misma fila: flecha horizontal
            flecha(slide, xa + W_ / 2, ya, xb - W_ / 2, yb)
        else:
            flecha(slide, xa, ya + H_ / 2, xb, yb - H_ / 2)
    for n, (cx, cy) in nodos.items():
        raiz, objetivo = n in ("Talento", "Distancia Madrid"), n == "Viabilidad empresarial"
        fill = NAVY if objetivo else (MID if raiz else TINT)
        caja(slide, cx - W_ / 2, cy - H_ / 2, W_, H_, fill=fill, line=None if (raiz or objetivo) else PALE, radio=0.25)
        texto(slide, cx - W_ / 2 + 0.05, cy - H_ / 2, W_ - 0.1, H_, [n], size=9.5, font=HEAD, bold=True,
              color=WHITE if (raiz or objetivo) else NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def barras_divergentes(slide, x, y, w, items, alto=0.62, vmin=-0.40, vmax=0.12, fmt="{:+.2f}"):
    """Barras horizontales con eje en cero dibujadas con formas (negativas en azul marino, positivas en pervinca)."""
    et_w = 1.95
    px = (w - et_w) / (vmax - vmin)
    x0 = x + et_w + (0 - vmin) * px
    for i, (lab, v) in enumerate(items):
        yy = y + i * alto
        texto(slide, x, yy, et_w - 0.1, alto * 0.7, [lab], size=10.5, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        bx, bw = (x0 + v * px, -v * px) if v < 0 else (x0, v * px)
        caja(slide, bx, yy + 0.08, bw, alto * 0.7 - 0.16, fill=NAVY if v < 0 else PERI, forma=MSO_SHAPE.RECTANGLE)
        lx = bx - 0.62 if v < 0 else bx + bw + 0.06
        texto(slide, lx, yy, 0.56, alto * 0.7, [fmt.format(v).replace("-", "−").replace(".", ",")], size=10.5, bold=True,
              color=INK, align=PP_ALIGN.RIGHT if v < 0 else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    eje = slide.shapes.add_connector(1, Inches(x0), Inches(y - 0.05), Inches(x0), Inches(y + len(items) * alto - 0.12))
    eje.line.color.rgb = GREY
    eje.line.width = Pt(1)


def entrada_salida(slide):
    """Dos tarjetas (entrada y salida) unidas por una flecha, para los ejemplos."""
    caja(slide, 0.79, 1.42, 3.55, 3.62)
    caja(slide, 4.87, 1.42, 4.34, 3.62)
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.42), Inches(3.05), Inches(0.38), Inches(0.36))
    a.fill.solid(); a.fill.fore_color.rgb = NAVY; a.line.fill.background()


def chip_regla(slide, x, y, tipo, txt, ancho_tipo=0.7):
    """Etiqueta de tipo (Filtro, Peso, Ámbito…) seguida de su texto."""
    filtro = tipo in ("Filtro", "Ámbito", "Política")
    caja(slide, x, y, ancho_tipo, 0.34, fill=NAVY if filtro else PERI, radio=0.5)
    texto(slide, x, y, ancho_tipo, 0.34, [tipo], size=9, color=WHITE if filtro else NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    texto(slide, x + ancho_tipo + 0.1, y, 3.15 - ancho_tipo, 0.34, [txt], size=10.5, anchor=MSO_ANCHOR.MIDDLE)


def barra_apilada(slide, x, y, w, partes):
    """Barra horizontal apilada con la aportación de cada parte a la puntuación (sobre 100)."""
    acum = 0.0
    for nom, v, color in partes:
        bw = w * v / 100
        caja(slide, x + acum, y, bw, 0.3, fill=color, forma=MSO_SHAPE.RECTANGLE)
        texto(slide, x + acum, y, bw, 0.3, [str(v)], size=9, bold=True, color=WHITE if color != PERI else NAVY,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        acum += bw
    lx = x  # leyenda debajo de la barra
    for nom, v, color in partes:
        caja(slide, lx, y + 0.42, 0.16, 0.16, fill=color, forma=MSO_SHAPE.RECTANGLE)
        texto(slide, lx + 0.22, y + 0.37, 1.1, 0.26, [nom], size=9.5, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
        lx += 1.25


def notas(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def nueva(prs, principal, resto=""):
    s = prs.slides.add_slide(prs.slide_layouts[4])  # TITLE_ONLY de la plantilla
    titulo(s, principal, resto)
    return s


def pie(slide, txt):
    texto(slide, 0.79, 5.18, 8.43, 0.25, [txt], size=8.5, color=GREY)


def dejar_solo(prs, orden):
    """Deja en la presentación solo las diapositivas de `orden`, en ese orden."""
    lst = prs.slides._sldIdLst
    por_slide = {}
    for sid in list(lst):
        por_slide[prs.part.related_part(sid.rId)] = sid
    quedan = {s.part for s in orden}
    for part, sid in por_slide.items():
        lst.remove(sid)
        if part not in quedan:
            prs.part.drop_rel(sid.rId)
    for s in orden:
        lst.append(por_slide[s.part])


def portada(prs, titulo_txt, subtitulo):
    s = prs.slides[0]
    lineas(por_id(s, 108), [titulo_txt])
    sub = por_id(s, 109)
    lineas(sub, subtitulo, size=11)
    sub.width = Inches(6.4)
    r0 = sub.text_frame.paragraphs[0].runs[0]  # el lema destaca sobre los autores
    r0.font.size, r0.font.bold = Pt(15), True
    return s


def cierre(prs, txt_extra):
    s = prs.slides[17]
    lineas(por_id(s, 734), ["Gracias"])
    lineas(por_id(s, 735), ["Gregorio García Velasco · UC3M", "Guillermo Franco Gimeno · UPM", txt_extra], size=13)
    # Se quitan los iconos de redes sociales y el aviso; el crédito a Slidesgo del layout se conserva (licencia gratuita).
    for sid in (736, 738, 739, 744):
        borrar(por_id(s, sid))
    return s


# ===================================================================================== PRESENTACIÓN 1
def presentacion_1():
    prs = Presentation(PLANTILLA)
    S = list(prs.slides)
    orden = []

    s = portada(prs, "Madrid 179",
                ["Dónde abrir la próxima oficina en la Comunidad de Madrid", AUTORES])
    notas(s, "Presentamos una herramienta de apoyo a la decisión con dos usuarios: empresas que buscan dónde abrir o trasladar "
             "oficinas, y administraciones que quieren descongestionar la capital y fijar actividad en el resto de la región.")
    orden.append(s)

    s = S[1]  # índice
    lineas(por_id(s, 115), ["Índice"])
    for sid, txt in zip((122, 123, 124, 125, 126, 127),
                        ("El problema", "El reto", "La solución", "Datos oficiales", "Uso empresarial", "Uso público")):
        lineas(por_id(s, sid), [txt])
    orden.append(s)

    # --- El problema (círculos de la plantilla) ---
    s = S[4]
    lineas(por_id(s, 200), ["La actividad se concentra en la capital"])
    for (a, b, c), (n, et, d) in zip(((206, 207, 208), (209, 210, 211), (212, 213, 214)),
                                     (("49,1 %", "Población", "vive en Madrid capital"),
                                      ("56,3 %", "Empresas", "de las unidades productivas"),
                                      ("64,3 %", "Empleo", "afiliado se ubica en la capital"))):
        lineas(por_id(s, a), [n], size=20); lineas(por_id(s, b), [et], size=13); lineas(por_id(s, c), [d], size=10.5)
    for sid in (201, 202):
        borrar(por_id(s, sid))
    texto(s, 0.79, 4.86, 8.43, 0.25,
          [[("Madrid capital ocupa el 7,5 % del territorio. ", True), ("142 de los 179 municipios tienen menos de 20.000 "
            "habitantes y reúnen solo el 8,8 % de la población.", False)]], size=10, align=PP_ALIGN.CENTER)
    pie(s, "Fuente: elaboración propia con datos del Instituto de Estadística de la Comunidad de Madrid (padrón 2025, colectivo empresarial 2025, afiliación 2025)")
    notas(s, "Cifras de nuestra tabla de 179 municipios. La capital ocupa el 7,5 % del territorio pero concentra la mitad de la "
             "población, más de la mitad de las unidades productivas y casi dos tercios del empleo afiliado localizado.")
    orden.append(s)

    # --- El reto (diagrama de Venn de la plantilla) ---
    s = S[2]
    lineas(por_id(s, 140), ["Dos públicos, una misma base de datos"])
    lineas(por_id(s, 144), ["Empresa"], size=10)
    lineas(por_id(s, 145), ["Adminis-", "tración"], size=10)
    lineas(por_id(s, 146), ["Datos", "CAM"], size=10.5)
    for (nom, des), (tn, td) in zip(
            [("Dónde ir", "Viabilidad del entorno"), ("Coste", "Coste, talento y transporte"), ("Apoyo", "Ayudas a las que optar"),
             ("Potencial", "Dónde hay margen"), ("Freno", "Factor que limita"), ("Política", "Escenarios simulados")],
            [(153, 154), (155, 156), (157, 158), (159, 160), (161, 162), (163, 164)]):
        lineas(por_id(s, tn), [nom], size=13)
        lineas(por_id(s, td), [des], size=10.5)
    notas(s, "A la izquierda, las preguntas de la empresa; a la derecha, las de la Administración. Ambas se responden con la "
             "misma tabla de 179 municipios construida con datos oficiales de la Comunidad de Madrid.")
    orden.append(s)

    # --- La solución (pasos de la plantilla) ---
    s = S[10]
    lineas(por_id(s, 465), ["La solución en cuatro pasos"])
    for sid, n in zip((466, 467, 468, 469), ("Paso 1", "Paso 2", "Paso 3", "Paso 4")):
        lineas(por_id(s, sid), [n], size=14)
    for (tn, td), (nom, des) in zip([(474, 475), (476, 477), (478, 479), (480, 481)],
                                    [("Filtros duros", "Fuera lo no negociable para la empresa"),
                                     ("GA²M", "Puntúa la viabilidad de cada municipio"),
                                     ("Red bayesiana", "Explica las causas y simula políticas"),
                                     ("Preferencias", "Transporte, ayudas y coste con peso visible")]):
        lineas(por_id(s, tn), [nom], size=14)
        lineas(por_id(s, td), [des], size=10.5)
    notas(s, "Primero se descarta lo imposible. Después el GA²M puntúa y la red bayesiana explica. Por último, las preferencias "
             "del usuario se suman con un peso explícito y calibrado, que se muestra por separado.")
    orden.append(s)

    # --- Datos oficiales ---
    s = nueva(prs, "Datos oficiales", "de la Comunidad")
    tiles = [("179", "municipios unidos por su código INE"), ("621", "tablas municipales del portal de datos abiertos de la Comunidad"),
             ("8 de 8", "variables del modelo obtenidas de fuentes de la Comunidad"), ("2.263", "conjuntos de datos catalogados en datos.comunidad.madrid")]
    w = (8.43 - 3 * 0.25) / 4
    for i, (n, d) in enumerate(tiles):
        x = 0.79 + i * (w + 0.25)
        caja(s, x, 1.45, w, 1.5)
        texto(s, x + 0.18, 1.6, w - 0.36, 0.5, [n], size=24, font=HEAD, bold=True, color=NAVY)
        texto(s, x + 0.18, 2.15, w - 0.36, 0.75, [d], size=10.5, color=GREY)
    texto(s, 0.79, 3.18, 8.43, 0.3, ["Organismos de la Comunidad de Madrid que alimentan el sistema"], size=13, font=HEAD, bold=True, color=NAVY)
    chips = ["Instituto de Estadística", "Portal de datos abiertos", "Consorcio Regional de Transportes",
             "Geoportal IDEM y cartografía", "D. G. de Economía e Industria", "Portal del Suelo"]
    cw = (8.43 - 2 * 0.2) / 3
    for i, c in enumerate(chips):
        x, y = 0.79 + (i % 3) * (cw + 0.2), 3.6 + (i // 3) * 0.55
        b = caja(s, x, y, cw, 0.42, fill=NAVY if i < 3 else MID, radio=0.5)
        texto(s, x, y, cw, 0.42, [c], size=11, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    pie(s, "Se complementan con INE, Dirección General del Catastro, Seguridad Social, BDNS y OpenStreetMap. Detalle en el documento de fuentes.")
    notas(s, "Las ocho variables del modelo, incluido el objetivo, se obtienen del portal de datos abiertos de la Comunidad o de "
             "organismos de la Comunidad. Algunas las produce el INE o el Catastro y las republica el Instituto de Estadística.")
    orden.append(s)

    # --- Cómo razona el sistema (conceptual, sin cifras del modelo) ---
    s = nueva(prs, "Cómo razona", "el sistema")
    texto(s, 0.79, 1.45, 2.65, 2.4, [
        [("Una red de causas ", True), ("que conecta el territorio con la viabilidad de una empresa.", False)],
        [("Explicable: ", True), ("cada recomendación dice qué factores la empujan y cuáles la frenan.", False)],
        [("Única para toda la región: ", True), ("el mismo razonamiento para los 179 municipios.", False)]], size=11, espacio=9)
    caja(s, 0.79, 3.95, 2.65, 1.05)
    texto(s, 0.95, 4.05, 2.35, 0.9, [[("Un GA²M", True, NAVY), (" ordena los municipios; la red bayesiana explica el porqué "
                                       "y permite simular políticas.", False)]], size=10.5)
    red_bayesiana_nativa(s)
    notas(s, "Las raíces son el talento y la distancia a Madrid. A través del coste, el dinamismo demográfico y la especialización "
             "llegan a la viabilidad empresarial. Así cada recomendación se puede explicar.")
    orden.append(s)

    # --- Explicabilidad nativa ---
    s = nueva(prs, "Explicable", "de forma nativa")
    for x in (0.79, 5.11):
        caja(s, x, 1.42, 4.1, 2.45)
    # GA²M: la puntuación es una suma de aportaciones
    circulo(s, 1.2, 1.8, 0.42, "1", fill=NAVY, size=12)
    texto(s, 1.55, 1.62, 3.2, 0.35, ["GA²M: cada variable suma o resta"], size=13, font=HEAD, bold=True, color=NAVY)
    texto(s, 1.0, 2.1, 3.7, 0.3, ["Puntuación = base + f(dinamismo) + f(servicios) + …"], size=9, color=GREY)
    barras_divergentes(s, 0.95, 2.45, 3.8, [("Dinamismo demográfico", 12), ("Especialización", 8), ("Distancia a Madrid", 3),
                                            ("Coste inmobiliario", -5)], alto=0.33, vmin=-11, vmax=16, fmt="{:+.0f}")
    # Red bayesiana: el porqué causal
    circulo(s, 5.52, 1.8, 0.42, "2", fill=MID, size=12)
    texto(s, 5.87, 1.62, 3.2, 0.35, ["Red bayesiana: el porqué causal"], size=13, font=HEAD, bold=True, color=NAVY)
    for i, (n, fill) in enumerate((("Coste", TINT), ("Dinamismo", TINT), ("Viabilidad", NAVY))):
        x = 5.35 + i * 1.3
        caja(s, x, 2.2, 1.05, 0.42, fill=WHITE if fill == TINT else NAVY, line=PALE if fill == TINT else None, radio=0.25)
        texto(s, x, 2.2, 1.05, 0.42, [n], size=9.5, font=HEAD, bold=True, color=NAVY if fill == TINT else WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 2:
            flecha(s, x + 1.05, 2.41, x + 1.3, 2.41, color=MID)
    texto(s, 5.33, 2.8, 3.7, 1.0, [
        [("Razona en los dos sentidos: ", True, NAVY), ("qué viabilidad cabe esperar y qué la explica.", False)],
        [("Permite intervenir: ", True, NAVY), ("¿y si baja el coste?", False)]], size=10.5, espacio=6)
    # Comparación con una caja negra
    caja(s, 0.79, 4.05, 4.1, 0.95, fill=WHITE, line=PALE)
    circulo(s, 1.13, 4.52, 0.38, "✗", fill=PALE, color=GREY, size=12, font=BODY)
    texto(s, 1.45, 4.13, 3.3, 0.8, [[("Otros modelos: caja negra", True, GREY)],
                                    "Se explican a posteriori, de forma aproximada, inestable y difícil de auditar."], size=10.5, color=GREY, espacio=3)
    caja(s, 5.11, 4.05, 4.1, 0.95, fill=NAVY)
    circulo(s, 5.45, 4.52, 0.38, "✓", fill=WHITE, color=NAVY, size=12, font=BODY)
    texto(s, 5.77, 4.13, 3.3, 0.8, [[("Nuestra propuesta: explicable por diseño", True, WHITE)],
                                    "Exacta, estable y verificable a mano: lo que se explica es lo que se calcula."], size=10.5, color=WHITE, espacio=3)
    pie(s, "Ejemplo ilustrativo: aportaciones simuladas.")
    notas(s, "Los dos modelos son de caja transparente. En el GA²M la puntuación es literalmente la suma de una aportación por "
             "variable, así que la explicación es exacta, no una aproximación posterior como SHAP o LIME. La red bayesiana añade "
             "el porqué causal: razona hacia delante y hacia atrás, y permite simular intervenciones. Para una administración "
             "pública esto significa decisiones que cualquiera puede revisar.")
    orden.append(s)

    # --- Ejemplo: empresa (datos simulados) ---
    s = nueva(prs, "Ejemplo:", "una empresa busca oficina")
    entrada_salida(s)
    texto(s, 0.97, 1.55, 3.2, 0.3, ["Entrada"], size=14, font=HEAD, bold=True, color=NAVY)
    texto(s, 0.97, 1.92, 3.2, 0.75, [[("«Estudio de diseño de 12 personas. Presupuesto medio, Cercanías imprescindible y, si puede "
                                       "ser, con ayudas.»", False, GREY)]], size=10.5)
    for i, (tipo, txt) in enumerate((("Filtro", "Cercanías a menos de 5 km"), ("Filtro", "Oferta de oficinas existente"),
                                     ("Peso", "Coste moderado · muy importante"), ("Peso", "Ayudas · poco importante"))):
        chip_regla(s, 0.97, 2.78 + i * 0.5, tipo, txt)
    texto(s, 5.05, 1.55, 3.9, 0.3, ["Salida: mejores opciones"], size=14, font=HEAD, bold=True, color=NAVY)
    for i, (nom, des, punt) in enumerate((("Municipio A", "Corredor del Henares · 45.000 hab.", "86"),
                                          ("Municipio B", "Sur metropolitano · 30.000 hab.", "79"),
                                          ("Municipio C", "Sierra Oeste · 12.000 hab.", "74"))):
        y = 1.98 + i * 0.52
        circulo(s, 5.25, y + 0.2, 0.36, str(i + 1), fill=NAVY if i == 0 else MID, size=11)
        texto(s, 5.55, y, 2.6, 0.24, [nom], size=11.5, font=HEAD, bold=True, color=NAVY)
        texto(s, 5.55, y + 0.24, 2.6, 0.22, [des], size=9.5, color=GREY)
        texto(s, 8.15, y + 0.02, 0.85, 0.4, [punt], size=17, font=HEAD, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
    texto(s, 5.05, 3.62, 3.9, 0.24, ["Por qué la opción 1"], size=10.5, font=HEAD, bold=True, color=NAVY)
    barra_apilada(s, 5.05, 3.92, 3.95, [("Viabilidad", 58, NAVY), ("Coste", 18, MID), ("Ayudas", 10, PERI)])
    texto(s, 5.05, 4.62, 3.95, 0.4, ["Buen dinamismo demográfico y especialización en servicios; estación a 2 km."], size=10, color=GREY)
    pie(s, "Ejemplo ilustrativo: los municipios y las cifras son simulados.")
    notas(s, "Ejemplo con datos simulados. La empresa describe lo que necesita en lenguaje natural; el sistema lo convierte en "
             "filtros (lo imprescindible) y pesos (lo deseable). La salida es un ranking con el porqué de cada opción: cuánto "
             "aporta la viabilidad del entorno y cuánto cada preferencia.")
    orden.append(s)

    # --- Ejemplo: administración (datos simulados) ---
    s = nueva(prs, "Ejemplo:", "evaluar una política")
    entrada_salida(s)
    texto(s, 0.97, 1.55, 3.2, 0.3, ["Entrada"], size=14, font=HEAD, bold=True, color=NAVY)
    for i, (tipo, txt) in enumerate((("Ámbito", "Municipios de la Sierra Norte"), ("Pregunta", "¿Qué frena al municipio X?"),
                                     ("Política", "Bonificación del alquiler de oficinas"), ("Hipótesis", "El coste baja un nivel (editable)"))):
        chip_regla(s, 0.97, 2.0 + i * 0.62, tipo, txt, ancho_tipo=0.85)
    texto(s, 5.05, 1.55, 3.9, 0.3, ["Salida: diagnóstico y escenario"], size=14, font=HEAD, bold=True, color=NAVY)
    barras_divergentes(s, 5.0, 1.98, 4.05, [("Dinamismo demográfico", -0.30), ("Nivel formativo", -0.10),
                                             ("Especialización", -0.05), ("Distancia a Madrid", 0.04)], alto=0.36)
    caja(s, 5.05, 3.5, 3.95, 1.05, fill=WHITE, line=PALE)
    texto(s, 5.2, 3.58, 3.7, 0.3, [[("Escenario: ", True, NAVY), ("puntuación de 22 a 29", False)]], size=11.5)
    texto(s, 5.2, 3.9, 3.7, 0.6, ["+7 puntos (entre +3 y +11). El freno principal es demográfico: conviene combinar la ayuda "
                                  "con políticas que atraigan población."], size=10, color=GREY)
    pie(s, "Ejemplo ilustrativo: el municipio y las cifras son simulados. Un escenario no es una predicción.")
    notas(s, "Ejemplo con datos simulados. El técnico define el ámbito y pregunta qué frena a un municipio: la ficha muestra qué "
             "factores restan y cuáles suman. Después prueba una política, ve la hipótesis con la que se traduce al modelo y "
             "obtiene un escenario con su margen de incertidumbre.")
    orden.append(s)

    # --- Arquitectura ---
    s = nueva(prs, "Arquitectura", "ligera y desacoplada")
    bloques = [("Aplicación de escritorio", "React y TypeScript con Tauri o Electron"), ("Mapa 3D", "MapLibre con Deck.gl, 179 municipios a 60 fps"),
               ("Motores de modelo", "Python, FastAPI, pgmpy e InterpretML en memoria"), ("Datos", "Parquet y GeoJSON; sin gestor de base de datos")]
    for i, (cab, cuerpo) in enumerate(bloques):
        x, y = 0.79 + (i % 2) * 4.32, 1.42 + (i // 2) * 1.82
        caja(s, x, y, 4.1, 1.6)
        circulo(s, x + 0.55, y + 0.8, 0.62, f"{i + 1}", fill=NAVY if i % 3 == 0 else MID, size=15)
        texto(s, x + 1.05, y + 0.35, 2.85, 0.4, [cab], size=14, font=HEAD, bold=True, color=NAVY)
        texto(s, x + 1.05, y + 0.8, 2.85, 0.7, [cuerpo], size=11, color=GREY)
    notas(s, "179 filas y datos anuales: no hace falta una base de datos pesada. Puntuar los 179 municipios lleva milisegundos.")
    orden.append(s)

    s = cierre(prs, "Madrid 179 · Localización inteligente de oficinas")
    orden.append(s)
    dejar_solo(prs, orden)
    prs.save("docs/presentaciones/1_Descripcion_y_problema.pptx")
    print("ok 1:", len(orden), "diapositivas")


# ===================================================================================== PRESENTACIÓN 3
def pantalla(prs, n, principal, resto, img, puntos, nota):
    s = nueva(prs, principal, resto)
    pic = s.shapes.add_picture(img, Inches(0.79), Inches(1.32), height=Inches(3.8))
    pic.line.color.rgb = PALE
    pic.line.width = Pt(1)
    texto(s, 0.79, 5.15, 4.3, 0.25, ["Prototipo de interfaz · las cifras de la pantalla son ilustrativas"], size=8.5, color=GREY)
    y = 1.42
    for i, (cab, des) in enumerate(puntos, start=1):
        circulo(s, 5.55, y + 0.19, 0.38, str(i), fill=NAVY, size=11)
        texto(s, 5.9, y, 3.3, 0.3, [cab], size=12.5, font=HEAD, bold=True, color=NAVY)
        texto(s, 5.9, y + 0.33, 3.3, 0.55, [des], size=10.5, color=GREY)
        y += 0.93
    notas(s, nota)
    return s


def presentacion_3():
    prs = Presentation(PLANTILLA)
    S = list(prs.slides)
    orden = []
    s = portada(prs, "Madrid 179",
                ["Interfaz: del perfil al mapa de oportunidades", AUTORES])
    notas(s, "Presentamos la interfaz: cómo llega cada usuario desde su perfil hasta una recomendación explicada.")
    orden.append(s)

    s = S[10]
    lineas(por_id(s, 465), ["Un recorrido en cuatro pasos"])
    for sid, n in zip((466, 467, 468, 469), ("Paso 1", "Paso 2", "Paso 3", "Paso 4")):
        lineas(por_id(s, sid), [n], size=14)
    for (tn, td), (nom, des) in zip([(474, 475), (476, 477), (478, 479), (480, 481)],
                                    [("Perfil", "Empresa o Administración"), ("Mi proyecto", "Necesidades que se vuelven filtros y pesos"),
                                     ("Mapa 3D", "La altura es la puntuación; preguntas que afinan"), ("Ficha e informe", "Desglose explicable y PDF")]):
        lineas(por_id(s, tn), [nom], size=14)
        lineas(por_id(s, td), [des], size=10.5)
    notas(s, "El recorrido es corto: elegir perfil, describir el proyecto, explorar el mapa y abrir la ficha del municipio.")
    orden.append(s)

    orden.append(pantalla(prs, 1, "Pantalla 1:", "inicio y perfil", "docs/presentaciones/img/interfaz-1.png", [
        ("Dos perfiles", "Empresas con acceso directo; Administración con acceso institucional."),
        ("Plantillas rápidas", "Startup, comercio, sede corporativa o logística como punto de partida."),
        ("Servicio público", "179 municipios y datos oficiales de la Comunidad de Madrid."),
        ("Ayuda siempre visible", "Atención 012 y guía rápida en la cabecera.")],
        "La entrada separa los dos perfiles desde el principio. El perfil de empresa es el de acceso directo; el de la "
        "Administración requiere acceso institucional."))
    orden.append(pantalla(prs, 2, "Pantalla 2:", "mi proyecto", "docs/presentaciones/img/interfaz-2.png", [
        ("Tipología y detalle libre", "Un LLM traduce la descripción en filtros duros y preferencias."),
        ("Presupuesto", "Filtro duro o peso de coste, según lo que indique el usuario."),
        ("Conexión prioritaria", "Preferencia de transporte: distancia a la estación más cercana."),
        ("Factores de ponderación", "Pesos explícitos y calibrados, mostrados por separado.")],
        "Cada control de esta pantalla tiene un efecto concreto en el cálculo: o descarta municipios o añade un peso visible."))
    orden.append(pantalla(prs, 3, "Pantalla 3:", "mapa de oportunidades", "docs/presentaciones/img/interfaz-3.png", [
        ("Altura del prisma", "Es la puntuación del municipio; los descartados quedan aplanados."),
        ("Afinar resultado", "La siguiente pregunta es la que más desempata a los candidatos."),
        ("Top de opciones", "Cada propuesta abre su ficha con el desglose de factores."),
        ("Informe PDF", "Recomendación, supuestos y fuentes en un documento.")],
        "El mapa responde al instante a cada respuesta. Las cifras de la captura son ilustrativas: en la versión final vendrán del modelo."))

    s = nueva(prs, "Cada respuesta", "tiene un efecto claro")
    tabla(s, 0.79, 1.45, 8.43, [["Respuesta del usuario", "Efecto en el cálculo", "Lectura para el usuario"],
                                ["Imprescindible", "Filtro duro", "Los municipios que no cumplen se descartan"],
                                ["Muy importante", "Peso completo", "Pesa tanto como el modelo"],
                                ["Poco importante", "Medio peso", "Pesa la mitad que el modelo"],
                                ["Me da igual", "Sin peso", "Solo cuenta la viabilidad"],
                                ["No lo sé", "Sin peso, se vuelve a preguntar", "La pregunta reaparece si desempata"]],
          [0.26, 0.34, 0.40], size=11.5, alto_fila=0.5)
    notas(s, "Las respuestas rápidas se traducen siempre igual, y la ficha muestra cuánto aporta cada preferencia.")
    orden.append(s)

    s = nueva(prs, "Perfil", "Administración")
    cw = (8.43 - 2 * 0.25) / 3
    for i, (cab, cuerpo, estado, fill) in enumerate([
            ("Diagnóstico", "Ámbito (región, comarca, tramo de población), mapa base y ficha por municipio con el factor que más lo frena.", "Disponible", NAVY),
            ("Simulador de políticas", "Política → palanca del modelo con hipótesis editable → intervención causal → escenario con intervalo.", "En desarrollo", MID),
            ("Informe", "Escenarios comparados con sus supuestos, intervalos y fuentes: escenario, no predicción.", "Diseñado", PERI)]):
        x = 0.79 + i * (cw + 0.25)
        caja(s, x, 1.42, cw, 2.45)
        circulo(s, x + 0.5, 1.95, 0.6, str(i + 1), fill=NAVY, size=15)
        texto(s, x + 0.25, 2.4, cw - 0.5, 0.35, [cab], size=14, font=HEAD, bold=True, color=NAVY)
        texto(s, x + 0.25, 2.85, cw - 0.5, 1.0, [cuerpo], size=10.5, color=GREY)
    texto(s, 0.79, 4.15, 8.43, 0.6, [[("Supuesto experto obligatorio: ", True, NAVY), ("si una palanca no tiene efecto observable en los datos, "
                                       "como una nueva estación, la herramienta exige declararlo y lo registra en el informe.", False)]], size=11.5)
    notas(s, "Pensado para la dirección general responsable del reequilibrio territorial, la consejería de economía y empleo, los "
             "ayuntamientos y las agencias de desarrollo local.")
    orden.append(s)

    s = nueva(prs, "El flujo", "de cada perfil")
    filas = [("Empresa", [("Perfil", 1), ("Describe", 0), ("Confirma", 1), ("Filtros", 1), ("Puntuación", 1), ("Mapa 3D", 1), ("Ficha e informe", 0)]),
             ("Administración", [("Perfil", 1), ("Ámbito", 1), ("Diagnóstico", 1), ("Política", 0), ("Hipótesis", 0), ("Escenario", 0), ("Informe", 0)])]
    for k, (nom, pasos) in enumerate(filas):
        y = 1.55 + k * 1.55
        texto(s, 0.79, y, 3, 0.3, [nom], size=13, font=HEAD, bold=True, color=NAVY)
        w, gap = 1.04, 0.19
        for i, (p, hecho) in enumerate(pasos):
            x = 0.79 + i * (w + gap)
            caja(s, x, y + 0.42, w, 0.62, fill=NAVY if k == 0 else MID, radio=0.2)
            texto(s, x + 0.04, y + 0.42, w - 0.08, 0.62, [p], size=10, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if i < len(pasos) - 1:
                a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w + 0.03), Inches(y + 0.66), Inches(0.13), Inches(0.14))
                a.fill.solid(); a.fill.fore_color.rgb = PERI; a.line.fill.background()
    notas(s, "Los dos perfiles comparten el mapa y los modelos; difieren en la pregunta: la empresa busca dónde ir, la "
             "Administración dónde actuar y con qué política.")
    orden.append(s)

    s = cierre(prs, "Madrid 179 · Interfaz")
    orden.append(s)
    dejar_solo(prs, orden)
    prs.save("docs/presentaciones/3_Interfaz.pptx")
    print("ok 3:", len(orden), "diapositivas")


if __name__ == "__main__":
    presentacion_1()
    presentacion_3()
